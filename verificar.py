"""
Auto-diagnóstico (não abre navegador): py verificar.py

Confere configuração dos clientes, ordem das perguntas, detecção de menções
(inclusive os falsos positivos de "lúmen"), extração e geração do Excel.
"""

import sys
import tempfile
from pathlib import Path

import analysis
import config
import storage

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

falhas = 0


def checar(cond, descricao):
    global falhas
    print(("  ✔ " if cond else "  ✘ ") + descricao)
    if not cond:
        falhas += 1


C = config.CLIENTES

print("Configuração")
checar(set(C) >= {"geracao", "kambalhota", "lumen"}, "clientes geracao, kambalhota e lumen existem")
checar(config.ORDEM_EXECUCAO == list(range(1, 26)) + [30, 26, 27, 28, 29], "ordem 1–25 → 30 → 26–29")
checar(config.ORDEM_EXECUCAO[-4:] == [26, 27, 28, 29], "26–29 são as 4 últimas")
pastas = [c["pasta_saida"] for c in C.values()]
checar(len(set(pastas)) == len(pastas), f"pastas de saída distintas: {pastas}")
checar(C["lumen"]["pasta_saida"] == "saida_lumen", "Lumen grava em saida_lumen/")
checar(C["geracao"]["pasta_saida"] == "saida" and C["kambalhota"]["pasta_saida"] == "saida_kambalhota",
       "geracao → saida/, kambalhota → saida_kambalhota/ (inalterados)")
checar(C["lumen"]["plataformas"] == ["chatgpt", "gemini", "perplexity"], "Lumen: chatgpt, gemini, perplexity")

print("\nPerguntas")
pt = config.pergunta_texto
checar(pt("geracao", 13) == "Escola de educação infantil no bairro Água Verde, em Curitiba.", "geracao P13 original")
checar(pt("geracao", 29) == "Geração do Saber ou outra escola do bairro Água Verde: qual é melhor?", "geracao P29 original")
checar(pt("geracao", 26) == "Vale a pena a Geração do Saber?", "geracao P26 original")
for n in (13, 14, 15, 29, 30):
    checar("Seminário" in pt("lumen", n) and "Água Verde" not in pt("lumen", n), f"lumen P{n} usa Seminário")
for n in (26, 27, 28, 29):
    checar("Escola Lumen" in pt("lumen", n) and "Geração" not in pt("lumen", n), f"lumen P{n}: {pt('lumen', n)}")
outras = [n for n in range(1, 31) if n not in (13, 14, 15, 26, 27, 28, 29, 30)]
checar(all(pt("lumen", n) == pt("geracao", n) == pt("kambalhota", n) for n in outras),
       "perguntas neutras idênticas entre os clientes")
checar(all("Lumen" not in pt("lumen", n) for n in range(1, 26)) and "Lumen" not in pt("lumen", 30),
       "nome da Lumen não aparece nas perguntas 1–25 e 30")

print("\nDetecção — Lumen (deve contar)")
L = C["lumen"]
positivos = [
    "Recomendo a Escola Lumen, no Seminário.",
    "O COLÉGIO LUMEN tem turmas pequenas.",
    "colegio lúmen é uma opção",
    "escola lumen",
    "Centro Educacional Lumen fica perto do Seminário.",
    "1. **Lumen** – educação infantil bilíngue\n2. Colégio Positivo",
    "- Lumen: escola com período integral",
    "Entre as opções no bairro Seminário estão a Lumen e o Positivo.",
]
for t in positivos:
    checar(analysis.detectar_mencao(t, L), repr(t))
print("Detecção — Lumen (NÃO deve contar)")
negativos = [
    "Uma lâmpada LED de 800 lúmens ilumina bem a sala de aula da escola.",
    "O fluxo luminoso é medido em lumen (lm).",
    "A escola usa projetores de 3000 Lumen nas salas.",
    "O documento Lumen Gentium, do Concílio Vaticano II, inspira escolas católicas.",
    "No lúmen intestinal ocorre a absorção.",
    "Colégio Positivo e Escola Internacional de Curitiba.",
    "Iluminação: Lumen é a unidade de fluxo luminoso.",
]
for t in negativos:
    checar(not analysis.detectar_mencao(t, L), repr(t))

print("\nDetecção — outros clientes")
K, G = C["kambalhota"], C["geracao"]
checar(analysis.detectar_mencao("Conheça o CE Kambalhota.", K), "CE Kambalhota")
checar(analysis.detectar_mencao("KAMBALHOTA é ótima", K), "KAMBALHOTA")
checar(analysis.detectar_mencao("centro educacional kambalhota", K), "centro educacional kambalhota")
checar(analysis.detectar_mencao("A GERACAO DO SABER fica no Água Verde", G), "GERACAO DO SABER sem acento")
checar(not analysis.detectar_mencao("Escola Lumen", G), "Lumen não conta como Geração do Saber")

print("\nExtração")
resp = ("Aqui estão algumas opções no Seminário:\n"
        "1. Colégio Positivo – tradicional, ensino forte.\n"
        "2. Escola Lumen – educação infantil com turmas pequenas e foco socioemocional.\n"
        "3. Maple Bear Curitiba – bilíngue.\n"
        "Fontes: https://www.escolalumen.com.br/ e guiaescolas.com.br/curitiba")
a = analysis.analisar(resp, [{"href": "https://www.google.com/url?q=https://reclameaqui.com.br/x"},
                             {"href": "https://chatgpt.com/c/123"}], L, "chatgpt")
checar(a["alvo_mencionado"] is True, "menção detectada")
checar(a["posicao"] == 2, f"posição na lista = 2 (obtido {a['posicao']})")
checar("Escola Lumen" in a["escolas"], f"escolas: {a['escolas']}")
checar("Escola Lumen" not in a["concorrentes"] and "Colégio Positivo" in a["concorrentes"]
       and "Maple Bear" in a["concorrentes"], f"concorrentes: {a['concorrentes']}")
checar("turmas pequenas" in a["descricao"], f"descrição: {a['descricao'][:80]}")
checar("https://reclameaqui.com.br/x" in a["fontes"], "link do Google desembrulhado")
checar(not any("chatgpt.com" in f for f in a["fontes"]), "link interno da plataforma ignorado")
checar(any("escolalumen.com.br" in f for f in a["fontes"]) and any("guiaescolas.com.br" in f for f in a["fontes"]),
       f"URLs/domínios no texto: {a['fontes']}")
checar(analysis.classificar_status("Não sei.") == "resposta_insuficiente", "resposta curta = insuficiente")
checar(analysis.analisar("Recomendo o Colégio Positivo e o Colégio Sion, ambos ótimos para crianças.", [], L)
       ["alvo_mencionado"] is False, "ausência real de menção = False")

print("\nExcel")
with tempfile.TemporaryDirectory() as tmp:
    log = Path(tmp) / "resultados.jsonl"
    base = {"cliente": "lumen", "data_hora": "2026-10-05 10:00:00", "observacoes": ""}
    storage.gravar_registro(log, {**base, "plataforma": "chatgpt", "numero": 1, "pergunta": pt("lumen", 1),
                                  "resposta": resp, "status": "ok", **a})
    storage.gravar_registro(log, {**base, "plataforma": "gemini", "numero": 1, "pergunta": pt("lumen", 1),
                                  "resposta": "", "status": "erro", "observacoes": "TimeoutError",
                                  "alvo_mencionado": None})
    with open(log, "a", encoding="utf-8") as f:
        f.write('{"linha corrompida"\n')
    regs = storage.ler_registros(log)
    checar(len(regs) == 2, "log lido, linha corrompida ignorada")
    destino = storage.exportar_excel("lumen", regs, Path(tmp) / "planilha.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(destino)
    checar(wb.sheetnames == ["Consultas", "Análise"], "abas Consultas e Análise")
    ws = wb["Consultas"]
    checar(ws.max_column == 13, "13 colunas na aba Consultas")
    valores = {ws.cell(r, 1).value: ws.cell(r, 6).value for r in range(2, ws.max_row + 1)}
    checar(valores.get("ChatGPT") == "Sim" and valores.get("Gemini") in ("", None),
           "Sim para menção; vazio (não 'Não') para falha de execução")

print()
if falhas:
    print(f"✘ {falhas} verificação(ões) falharam.")
    sys.exit(1)
print("✔ Todas as verificações passaram.")
