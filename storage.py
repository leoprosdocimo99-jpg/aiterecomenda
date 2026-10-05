"""
Armazenamento dos resultados.

- resultados.jsonl: log append-only, 1 linha por tentativa, gravado e
  sincronizado com o disco imediatamente (fonte da verdade; nunca é apagado).
- pesquisa_visibilidade.xlsx: regenerado a partir do log após cada consulta
  (aba Consultas + aba Análise). Se o Excel estiver aberto e bloqueado,
  grava uma cópia com outro nome em vez de falhar.
"""

import json
import os
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import config

STATUS_CONCLUIDOS = {"ok", "resposta_insuficiente"}


# ---------------------------------------------------------------------------
# Log JSONL
# ---------------------------------------------------------------------------
def gravar_registro(caminho: Path, registro: dict) -> None:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with open(caminho, "a", encoding="utf-8") as f:
        f.write(json.dumps(registro, ensure_ascii=False) + "\n")
        f.flush()
        os.fsync(f.fileno())


def ler_registros(caminho: Path) -> list:
    """Lê o log; linhas corrompidas (ex.: queda de energia no meio da escrita) são ignoradas."""
    if not caminho.exists():
        return []
    registros = []
    with open(caminho, encoding="utf-8") as f:
        for n, linha in enumerate(f, 1):
            linha = linha.strip()
            if not linha:
                continue
            try:
                registros.append(json.loads(linha))
            except json.JSONDecodeError:
                print(f"  [aviso] linha {n} de {caminho.name} ilegível — ignorada")
    return registros


def ultimos_por_consulta(registros: list) -> dict:
    """{(plataforma, numero): último registro} — tentativa mais recente vale."""
    ultimos = {}
    for r in registros:
        ultimos[(r.get("plataforma"), r.get("numero"))] = r
    return ultimos


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _sim_nao(r):
    if r.get("status") not in STATUS_CONCLUIDOS:
        return ""  # falha de execução: nem Sim nem Não
    return "Sim" if r.get("alvo_mencionado") else "Não"


def _curtos(cliente):
    c = cliente.get("nome_curto") or []
    return ", ".join(f'"{x}"' for x in ([c] if isinstance(c, str) else c))


def _lista(v):
    return "; ".join(v) if isinstance(v, list) else (v or "")


def exportar_excel(cliente_id: str, registros: list, destino: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cliente = config.CLIENTES[cliente_id]
    nome = cliente["nome"]
    ultimos = ultimos_por_consulta(registros)
    ordem_q = {n: i for i, n in enumerate(config.ORDEM_EXECUCAO)}
    ordem_p = {p: i for i, p in enumerate(config.ORDEM_PLATAFORMAS)}
    linhas = sorted(ultimos.values(), key=lambda r: (
        ordem_p.get(r.get("plataforma"), 99), ordem_q.get(r.get("numero"), 99)))

    wb = Workbook()
    negrito = Font(bold=True)
    cab_fill = PatternFill("solid", fgColor="DDEBF7")
    quebra = Alignment(wrap_text=True, vertical="top")

    # ---------------- Aba Consultas ----------------
    ws = wb.active
    ws.title = "Consultas"
    colunas = [
        ("Plataforma", 12), ("Nº pergunta", 10), ("Pergunta original", 40),
        ("Resposta integral da IA", 80), ("Escolas mencionadas", 40),
        (f"{nome} mencionada (Sim/Não)", 16), (f"Posição da {nome} em lista ordenada", 14),
        ("Concorrentes citados", 40), (f"Descrição atribuída à {nome}", 50),
        ("Fontes e URLs mencionadas", 50), ("Data e horário da consulta", 19),
        ("Status da execução", 18), ("Observações e erros", 40),
    ]
    ws.append([c for c, _ in colunas])
    for i, (_, larg) in enumerate(colunas, 1):
        ws.column_dimensions[get_column_letter(i)].width = larg
        ws.cell(1, i).font = negrito
        ws.cell(1, i).fill = cab_fill
    for r in linhas:
        ws.append([
            config.CONFIG_PLATAFORMAS.get(r.get("plataforma"), {}).get("nome", r.get("plataforma")),
            r.get("numero"),
            r.get("pergunta"),
            (r.get("resposta") or "")[:32000],  # limite de célula do Excel
            _lista(r.get("escolas")),
            _sim_nao(r),
            r.get("posicao") if r.get("posicao") is not None else "",
            _lista(r.get("concorrentes")),
            r.get("descricao") or "",
            "\n".join(r.get("fontes") or []),
            r.get("data_hora"),
            r.get("status"),
            r.get("observacoes") or "",
        ])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = quebra
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ---------------- Aba Análise ----------------
    wa = wb.create_sheet("Análise")
    wa.column_dimensions["A"].width = 45
    for col in "BCDEFG":
        wa.column_dimensions[col].width = 18

    def titulo(texto):
        wa.append([])
        wa.append([texto])
        wa.cell(wa.max_row, 1).font = Font(bold=True, size=12)

    def cabecalho(valores):
        wa.append(valores)
        for i in range(1, len(valores) + 1):
            wa.cell(wa.max_row, i).font = negrito
            wa.cell(wa.max_row, i).fill = cab_fill

    wa.append([f"Pesquisa de visibilidade em IAs — escola-alvo: {nome} "
               f"(bairro {cliente['bairro']}, {config.CIDADE})"])
    wa.cell(1, 1).font = Font(bold=True, size=14)
    wa.append([f"Grafias aceitas: {', '.join(cliente['variantes'])}"
               + (f"; e {_curtos(cliente)} quando usado(s) como nome de escola"
                  if cliente.get("nome_curto") else "")])
    wa.append([f"Gerado em {datetime.now():%d/%m/%Y %H:%M}"])

    plataformas = [p for p in config.ORDEM_PLATAFORMAS if p in cliente["plataformas"]
                   or any(r.get("plataforma") == p for r in linhas)]

    titulo("1. Menções por plataforma")
    cabecalho(["Plataforma", "Consultas registradas", "Respondidas (ok)",
               "Resposta insuficiente", "Falhas de execução", f"Menções à {nome}",
               "% de menção (sobre respondidas)"])
    for p in plataformas:
        rs = [r for r in linhas if r.get("plataforma") == p]
        ok = [r for r in rs if r.get("status") == "ok"]
        insuf = [r for r in rs if r.get("status") == "resposta_insuficiente"]
        falhas = [r for r in rs if r.get("status") not in STATUS_CONCLUIDOS]
        respondidas = ok + insuf
        mencoes = [r for r in respondidas if r.get("alvo_mencionado")]
        pct = f"{100 * len(mencoes) / len(respondidas):.1f}%" if respondidas else "—"
        wa.append([config.CONFIG_PLATAFORMAS[p]["nome"], len(rs), len(ok), len(insuf),
                   len(falhas), len(mencoes), pct])

    titulo(f"2. Perguntas que geraram menção à {nome}")
    cabecalho(["Pergunta", "Nº", "Plataforma", "Posição em lista", "Descrição atribuída"])
    for r in linhas:
        if r.get("status") in STATUS_CONCLUIDOS and r.get("alvo_mencionado"):
            wa.append([r.get("pergunta"), r.get("numero"),
                       config.CONFIG_PLATAFORMAS[r["plataforma"]]["nome"],
                       r.get("posicao") or "", (r.get("descricao") or "")[:500]])

    titulo("3. Concorrentes mais recorrentes")
    cabecalho(["Escola", "Total de respostas"] + [config.CONFIG_PLATAFORMAS[p]["nome"] for p in plataformas])
    total, por_plat = Counter(), defaultdict(Counter)
    for r in linhas:
        if r.get("status") in STATUS_CONCLUIDOS:
            for e in set(r.get("concorrentes") or []):
                total[e] += 1
                por_plat[r["plataforma"]][e] += 1
    for escola, n in total.most_common(30):
        wa.append([escola, n] + [por_plat[p][escola] for p in plataformas])

    titulo(f"4. Comparação entre plataformas (menção à {nome} por pergunta)")
    cabecalho(["Pergunta", "Nº"] + [config.CONFIG_PLATAFORMAS[p]["nome"] for p in plataformas])
    for n in config.ORDEM_EXECUCAO:
        linha = [config.pergunta_texto(cliente_id, n), n]
        for p in plataformas:
            r = ultimos.get((p, n))
            if r is None:
                linha.append("não executada")
            elif r.get("status") not in STATUS_CONCLUIDOS:
                linha.append(f"falha ({r.get('status')})")
            elif r.get("alvo_mencionado"):
                linha.append("Sim" + (f" (#{r['posicao']})" if r.get("posicao") else ""))
            else:
                linha.append("Não" + (" (insuf.)" if r.get("status") == "resposta_insuficiente" else ""))
        wa.append(linha)

    titulo("5. Nota metodológica")
    for t in [
        "Cada pergunta foi feita em uma conversa nova, sem histórico, na ordem 1–25 → 30 → 26–29.",
        "As perguntas 1–25 e 30 são neutras; o nome da escola só aparece nas perguntas 26–29.",
        "\"Não\" = a IA respondeu e não citou a escola (ausência real de menção).",
        "\"falha\" (erro, erro_login) = problema técnico; não conta como presença nem ausência.",
        "\"resposta_insuficiente\" = a IA respondeu, mas sem conteúdo aproveitável.",
        "Percentuais calculados sobre consultas respondidas (ok + insuficiente).",
        "Escolas e descrições são extraídas automaticamente do texto da IA; nada é inventado. "
        "Recomenda-se conferência manual da coluna \"Escolas mencionadas\".",
        "Vale a tentativa mais recente de cada consulta; o histórico completo está em resultados.jsonl.",
    ]:
        wa.append([t])

    destino.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(destino)
        return destino
    except PermissionError:
        alternativo = destino.with_name(f"{destino.stem}_{datetime.now():%Y%m%d_%H%M%S}{destino.suffix}")
        wb.save(alternativo)
        print(f"  [aviso] {destino.name} está aberto no Excel; salvo como {alternativo.name}")
        return alternativo
